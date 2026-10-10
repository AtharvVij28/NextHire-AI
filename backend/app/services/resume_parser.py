from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.config import MAX_RESUME_CHARS


class ResumeParseError(Exception):
    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


def extract_resume_text(data: bytes, extension: str) -> str:
    if extension == ".pdf":
        text = _extract_pdf(data)
    elif extension == ".docx":
        text = _extract_docx(data)
    else:
        raise ResumeParseError("Only PDF and DOCX resumes are supported.")

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise ResumeParseError(
            "Could not read text from this file. Scanned image PDFs are not supported — upload a text-based PDF or DOCX."
        )
    if len(text) > MAX_RESUME_CHARS:
        text = text[:MAX_RESUME_CHARS]
    return text


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise ResumeParseError("This PDF could not be opened.") from exc

    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise ResumeParseError("This DOCX file could not be opened.") from exc

    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)
